import argparse
import json
import sys
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
DEFAULT_FONT = "Microsoft YaHei"
TITLE_FONT = "Microsoft YaHei UI"
DEFAULT_OUTPUT = "generated/student-ppt-pet/My_Academic_Presentation.pptx"
DEFAULT_STATE = "state.json"
DEFAULT_TEMPLATE_KEY = "academic-multi"
if getattr(sys, "frozen", False):
    EXEC_ROOT = Path(sys.executable).resolve().parent
    SCRIPT_ROOT = Path(getattr(sys, "_MEIPASS", EXEC_ROOT))
    SKILL_ROOT = EXEC_ROOT.parent.parent
    TEMPLATE_DIR_CANDIDATES = [
        EXEC_ROOT / "assets" / "templates",
        SKILL_ROOT / "assets" / "templates",
        SCRIPT_ROOT / "assets" / "templates",
    ]
else:
    SCRIPT_ROOT = Path(__file__).resolve().parent
    SKILL_ROOT = SCRIPT_ROOT.parent
    TEMPLATE_DIR_CANDIDATES = [SKILL_ROOT / "assets" / "templates"]

TEMPLATE_DIR = next(
    (path for path in TEMPLATE_DIR_CANDIDATES if path.exists()),
    TEMPLATE_DIR_CANDIDATES[0],
)

THEME = {
    "navy": "0F2747",
    "blue": "1F5EFF",
    "sky": "EAF2FF",
    "cyan": "4DA8FF",
    "ink": "1E2A39",
    "gray": "5E6B78",
    "light": "F7F9FC",
    "white": "FFFFFF",
    "border": "D8E1F0",
    "accent": "FF8A3D",
    "green": "1FA971",
    "panel": "17365F",
    "panel_alt": "214C99",
}
ACTIVE_TEMPLATE = {
    "key": DEFAULT_TEMPLATE_KEY,
    "display_name": "Academic Multi",
    "fonts": {"body": DEFAULT_FONT, "title": TITLE_FONT},
    "theme": deepcopy(THEME),
    "default_variants": {
        "title": "hero-band",
        "agenda": "grid-cards",
        "content": "sidebar-focus",
        "process": "step-cards",
        "results": "metrics-chart",
        "summary": "takeaway-panel",
    },
    "variants": {
        "title": ["hero-band", "center-spotlight"],
        "agenda": ["grid-cards", "stacked-list"],
        "content": ["sidebar-focus", "split-panel"],
        "process": ["step-cards", "timeline-ribbon"],
        "results": ["metrics-chart", "insight-grid"],
        "summary": ["takeaway-panel", "closing-quote"],
    },
}


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def load_template_registry() -> dict:
    templates = {}
    if TEMPLATE_DIR.exists():
        for path in sorted(TEMPLATE_DIR.glob("*.json")):
            data = json.loads(path.read_text(encoding="utf-8"))
            key = str(data.get("key", path.stem)).strip() or path.stem
            data["_path"] = str(path)
            templates[key] = data
    if DEFAULT_TEMPLATE_KEY not in templates:
        templates[DEFAULT_TEMPLATE_KEY] = deepcopy(ACTIVE_TEMPLATE)
        templates[DEFAULT_TEMPLATE_KEY]["_path"] = ""
    return templates


TEMPLATE_REGISTRY = load_template_registry()


def set_active_template(template_key: str | None) -> dict:
    global ACTIVE_TEMPLATE
    global THEME
    global DEFAULT_FONT
    global TITLE_FONT

    chosen_key = str(template_key or DEFAULT_TEMPLATE_KEY).strip() or DEFAULT_TEMPLATE_KEY
    ACTIVE_TEMPLATE = deepcopy(TEMPLATE_REGISTRY.get(chosen_key, TEMPLATE_REGISTRY[DEFAULT_TEMPLATE_KEY]))
    THEME = deepcopy(ACTIVE_TEMPLATE.get("theme", {}))
    DEFAULT_FONT = str(ACTIVE_TEMPLATE.get("fonts", {}).get("body", "Microsoft YaHei"))
    TITLE_FONT = str(ACTIVE_TEMPLATE.get("fonts", {}).get("title", DEFAULT_FONT))
    return ACTIVE_TEMPLATE


def get_allowed_variants(layout: str) -> list[str]:
    return list(ACTIVE_TEMPLATE.get("variants", {}).get(layout, []))


def get_default_variant(layout: str) -> str:
    return str(ACTIVE_TEMPLATE.get("default_variants", {}).get(layout, layout)).strip() or layout


def get_known_layouts() -> list[str]:
    return list(ACTIVE_TEMPLATE.get("variants", {}).keys())


def get_variant_layout_map() -> dict[str, str]:
    mapping = {}
    for layout, variants in ACTIVE_TEMPLATE.get("variants", {}).items():
        for variant in variants:
            mapping[str(variant).strip()] = str(layout).strip()
    return mapping


def resolve_template_key(
    explicit_template_key: str | None, payload: dict | None = None
) -> str:
    if explicit_template_key:
        value = str(explicit_template_key).strip()
        if value:
            return value

    if isinstance(payload, dict):
        for key in ("template_key", "template"):
            value = str(payload.get(key, "")).strip()
            if value:
                return value

    return DEFAULT_TEMPLATE_KEY


def rgb(hex_color: str) -> RGBColor:
    hex_color = str(hex_color).replace("#", "")
    return RGBColor(
        int(hex_color[0:2], 16),
        int(hex_color[2:4], 16),
        int(hex_color[4:6], 16),
    )


def normalize_lines(items):
    if items is None:
        return []
    if isinstance(items, str):
        text = items.strip()
        return [text] if text else []
    if not isinstance(items, list):
        return [str(items).strip()] if str(items).strip() else []
    return [str(item).strip() for item in items if str(item).strip()]


def ensure_slide_size(prs: Presentation) -> None:
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H


def set_background(slide, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_box(slide, left, top, width, height, fill_color, line_color=None, radius=True):
    shape_type = (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        if radius
        else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill_color)
    shape.line.color.rgb = rgb(line_color or fill_color)
    return shape


def add_line(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.color.rgb = rgb(color)
    return shape


def style_text_frame(
    text_frame,
    *,
    font_name=DEFAULT_FONT,
    font_size=20,
    color="1E2A39",
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    text_frame.vertical_anchor = valign
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = align
        for run in paragraph.runs:
            run.font.name = font_name
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    font_name=DEFAULT_FONT,
    font_size=18,
    color="1E2A39",
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    textbox = slide.shapes.add_textbox(left, top, width, height)
    tf = textbox.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = str(text)
    style_text_frame(
        tf,
        font_name=font_name,
        font_size=font_size,
        color=color,
        bold=bold,
        align=align,
        valign=valign,
    )
    return textbox


def add_bullets(
    slide,
    left,
    top,
    width,
    height,
    bullets,
    *,
    font_size=18,
    color="1E2A39",
    level=0,
    max_items=6,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True

    items = normalize_lines(bullets)[:max_items] or ["Add content"]
    first = tf.paragraphs[0]
    first.text = items[0]
    first.level = level
    first.space_after = Pt(10)
    style_text_frame(tf, font_size=font_size, color=color)

    for item in items[1:]:
        p = tf.add_paragraph()
        p.text = item
        p.level = level
        p.space_after = Pt(10)
        for run in p.runs:
            run.font.name = DEFAULT_FONT
            run.font.size = Pt(font_size)
            run.font.color.rgb = rgb(color)
    return box


def add_footer(slide, page_num: int):
    add_line(
        slide, Inches(0.6), Inches(7.05), Inches(12.1), Inches(0.03), THEME["border"]
    )
    add_text(
        slide,
        Inches(0.7),
        Inches(7.08),
        Inches(4.5),
        Inches(0.25),
        "Student PPT Pet | Academic Presentation",
        font_size=9,
        color=THEME["gray"],
    )
    badge = add_box(
        slide, Inches(12.05), Inches(6.93), Inches(0.55), Inches(0.34), THEME["navy"]
    )
    add_text(
        slide,
        badge.left,
        badge.top + Inches(0.01),
        badge.width,
        badge.height,
        str(page_num),
        font_name=TITLE_FONT,
        font_size=12,
        color=THEME["white"],
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def infer_layout(slide_data: dict, *, allow_explicit: bool = True) -> str:
    layout = str(slide_data.get("layout", "")).lower().strip()
    if allow_explicit and layout in get_known_layouts():
        return layout

    slide_type = str(slide_data.get("type", "content")).lower()
    if slide_type == "title":
        return "title"

    title = str(slide_data.get("title", "")).lower()
    if any(k in title for k in ("agenda", "目录")):
        return "agenda"
    if any(k in title for k in ("method", "process", "flow", "architecture", "方法", "流程", "架构", "实验设置")):
        return "process"
    if any(k in title for k in ("result", "analysis", "data", "benchmark", "结果", "分析", "数据", "对比")):
        return "results"
    if any(k in title for k in ("summary", "conclusion", "thanks", "总结", "结论", "感谢")):
        return "summary"
    return "content"


def canonicalize_layout_variant(slide_data: dict) -> tuple[str, str]:
    raw_layout = str(slide_data.get("layout", "")).lower().strip()
    raw_variant = str(slide_data.get("variant", "")).strip()
    known_layouts = set(get_known_layouts())
    variant_layout_map = get_variant_layout_map()

    layout = raw_layout
    variant = raw_variant

    if layout in variant_layout_map:
        aliased_variant = layout
        layout = variant_layout_map[aliased_variant]
        if not variant:
            variant = aliased_variant

    if variant in variant_layout_map:
        layout = variant_layout_map[variant]

    if layout not in known_layouts:
        layout = infer_layout(slide_data, allow_explicit=False)

    if variant in variant_layout_map:
        layout = variant_layout_map[variant]
    elif variant not in get_allowed_variants(layout):
        variant = get_default_variant(layout)

    if not variant:
        variant = get_default_variant(layout)

    return layout, variant


def normalize_slide(slide_data: dict, slide_number: int) -> dict:
    slide = deepcopy(slide_data if isinstance(slide_data, dict) else {})
    slide["slide_number"] = slide_number
    slide["type"] = str(slide.get("type", "content")).lower()
    slide["layout"], variant = canonicalize_layout_variant(slide)
    slide["variant"] = variant
    slide["title"] = str(slide.get("title", "")).strip()
    slide["content"] = normalize_lines(slide.get("content", []))

    metrics = slide.get("metrics", [])
    if isinstance(metrics, list):
        normalized_metrics = []
        for metric in metrics:
            if not isinstance(metric, dict):
                continue
            label = str(metric.get("label", "")).strip()
            value = str(metric.get("value", "")).strip()
            if label or value:
                normalized_metrics.append({"label": label, "value": value})
        slide["metrics"] = normalized_metrics
    else:
        slide["metrics"] = []

    chart = slide.get("chart", {})
    if isinstance(chart, dict):
        slide["chart"] = {
            "series_name": str(chart.get("series_name", "Series")).strip() or "Series",
            "categories": normalize_lines(chart.get("categories", [])),
            "values": list(chart.get("values", []))
            if isinstance(chart.get("values", []), list)
            else [],
        }
    else:
        slide["chart"] = {}
    return slide


def normalize_plan(plan: dict, template_key: str | None = None) -> dict:
    if not isinstance(plan, dict):
        raise ValueError("plan must be a JSON object")
    chosen_template_key = resolve_template_key(template_key, plan)
    set_active_template(chosen_template_key)

    slides = plan.get("slides", [])
    if not isinstance(slides, list) or not slides:
        raise ValueError("slides must be a non-empty list")

    normalized = {
        "title": str(plan.get("title", "Academic Presentation")).strip()
        or "Academic Presentation",
        "author": str(plan.get("author", "Unknown Author")).strip() or "Unknown Author",
        "template_key": ACTIVE_TEMPLATE["key"],
        "slides": [],
    }
    for idx, slide_data in enumerate(slides, start=1):
        normalized["slides"].append(normalize_slide(slide_data, idx))
    return normalized


def extract_plan_from_state(state: dict) -> dict:
    if not isinstance(state, dict):
        raise ValueError("state must be a JSON object")
    if isinstance(state.get("plan"), dict):
        return normalize_plan(state["plan"], state.get("template_key"))
    return normalize_plan(state, state.get("template_key"))


def restyle_plan_to_template_defaults(plan: dict, template_key: str) -> dict:
    normalized = normalize_plan(plan, template_key)
    restyled = deepcopy(normalized)
    set_active_template(template_key)
    for slide in restyled["slides"]:
        layout = slide["layout"]
        slide["variant"] = get_default_variant(layout)
    return normalize_plan(restyled, template_key)


def make_state(plan: dict, *, template_path: str, output_path: str, template_key: str) -> dict:
    return {
        "version": 1,
        "engine": {"name": "student-ppt-pet", "mode": "local-exe"},
        "updated_at": now_iso(),
        "template_key": template_key,
        "template_name": ACTIVE_TEMPLATE.get("display_name", template_key),
        "template_path": template_path,
        "output_path": output_path,
        "plan": plan,
    }


def save_json(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def load_json(path: Path) -> dict:
    if not path.exists():
        raise FileNotFoundError(f"json file not found: {path}")
    return json.loads(path.read_text(encoding="utf-8-sig"))


class RelaxedJsonParser:
    def __init__(self, text: str):
        self.text = text.strip()
        self.length = len(self.text)
        self.index = 0

    def parse(self):
        value = self.parse_value()
        self.skip_ws()
        if self.index != self.length:
            raise ValueError(f"unexpected trailing content near index {self.index}")
        return value

    def current(self) -> str:
        if self.index >= self.length:
            return ""
        return self.text[self.index]

    def consume(self, expected: str | None = None) -> str:
        if self.index >= self.length:
            raise ValueError("unexpected end of input")
        char = self.text[self.index]
        if expected is not None and char != expected:
            raise ValueError(f"expected '{expected}' at index {self.index}, got '{char}'")
        self.index += 1
        return char

    def skip_ws(self) -> None:
        while self.index < self.length and self.text[self.index].isspace():
            self.index += 1

    def parse_value(self):
        self.skip_ws()
        char = self.current()
        if char == "{":
            return self.parse_object()
        if char == "[":
            return self.parse_array()
        if char in ("'", '"'):
            return self.parse_string()
        return self.parse_literal()

    def parse_object(self) -> dict:
        obj = {}
        self.consume("{")
        self.skip_ws()
        if self.current() == "}":
            self.consume("}")
            return obj
        while True:
            key = self.parse_key()
            self.skip_ws()
            self.consume(":")
            obj[key] = self.parse_value()
            self.skip_ws()
            char = self.current()
            if char == ",":
                self.consume(",")
                self.skip_ws()
                continue
            if char == "}":
                self.consume("}")
                return obj
            raise ValueError(f"expected ',' or '}}' at index {self.index}")

    def parse_array(self) -> list:
        arr = []
        self.consume("[")
        self.skip_ws()
        if self.current() == "]":
            self.consume("]")
            return arr
        while True:
            arr.append(self.parse_value())
            self.skip_ws()
            char = self.current()
            if char == ",":
                self.consume(",")
                self.skip_ws()
                continue
            if char == "]":
                self.consume("]")
                return arr
            raise ValueError(f"expected ',' or ']' at index {self.index}")

    def parse_key(self) -> str:
        self.skip_ws()
        char = self.current()
        if char in ("'", '"'):
            return self.parse_string()
        return self.parse_bareword(stop_chars={":", "}", ","})

    def parse_string(self) -> str:
        quote = self.consume()
        chunks = []
        while True:
            if self.index >= self.length:
                raise ValueError("unterminated string")
            char = self.consume()
            if char == "\\":
                if self.index >= self.length:
                    raise ValueError("unterminated escape sequence")
                nxt = self.consume()
                escape_map = {
                    '"': '"',
                    "'": "'",
                    "\\": "\\",
                    "/": "/",
                    "b": "\b",
                    "f": "\f",
                    "n": "\n",
                    "r": "\r",
                    "t": "\t",
                }
                chunks.append(escape_map.get(nxt, nxt))
                continue
            if char == quote:
                return "".join(chunks)
            chunks.append(char)

    def parse_bareword(self, *, stop_chars: set[str]) -> str:
        start = self.index
        while self.index < self.length and self.text[self.index] not in stop_chars:
            self.index += 1
        value = self.text[start:self.index].strip()
        if not value:
            raise ValueError(f"empty token near index {start}")
        return value

    def parse_literal(self):
        token = self.parse_bareword(stop_chars={",", "}", "]"})
        lowered = token.lower()
        if lowered == "true":
            return True
        if lowered == "false":
            return False
        if lowered == "null":
            return None
        try:
            if any(ch in token for ch in (".", "e", "E")):
                return float(token)
            return int(token)
        except ValueError:
            return token


def parse_json_payload(text: str):
    payload = str(text or "").strip()
    if not payload:
        raise ValueError("empty JSON payload")
    try:
        return json.loads(payload)
    except json.JSONDecodeError:
        return RelaxedJsonParser(payload).parse()


def add_cover_slide(prs: Presentation, slide_data: dict):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    variant = slide_data.get("variant", get_default_variant("title"))
    subtitle = "\n".join(normalize_lines(slide_data.get("content", [])))
    if variant == "center-spotlight":
        set_background(slide, THEME["light"])
        add_box(slide, Inches(0.6), Inches(0.65), Inches(12.1), Inches(0.22), THEME["accent"], radius=False)
        add_box(slide, Inches(1.0), Inches(1.45), Inches(11.2), Inches(4.6), THEME["white"], THEME["border"])
        add_box(slide, Inches(9.85), Inches(1.15), Inches(1.55), Inches(1.55), THEME["cyan"])
        add_box(slide, Inches(10.55), Inches(4.9), Inches(1.0), Inches(1.0), THEME["green"])
        add_text(
            slide,
            Inches(1.55),
            Inches(2.1),
            Inches(9.3),
            Inches(0.9),
            slide_data.get("title", "Academic Presentation"),
            font_name=TITLE_FONT,
            font_size=28,
            color=THEME["navy"],
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            Inches(2.0),
            Inches(3.1),
            Inches(8.4),
            Inches(1.25),
            subtitle or "Course report / academic defense / research update",
            font_size=16,
            color=THEME["gray"],
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            Inches(4.3),
            Inches(5.3),
            Inches(4.5),
            Inches(0.35),
            ACTIVE_TEMPLATE.get("display_name", "Student PPT Pet"),
            font_size=11,
            color=THEME["blue"],
            align=PP_ALIGN.CENTER,
        )
        return slide

    set_background(slide, THEME["navy"])
    add_box(slide, Inches(0.8), Inches(0.75), Inches(0.18), Inches(5.8), THEME["accent"], radius=False)
    add_box(slide, Inches(9.7), Inches(0.8), Inches(2.2), Inches(2.2), THEME["blue"])
    add_box(slide, Inches(10.35), Inches(1.45), Inches(1.2), Inches(1.2), THEME["cyan"])
    add_box(slide, Inches(8.85), Inches(4.95), Inches(2.7), Inches(1.05), THEME["panel_alt"])
    add_text(
        slide,
        Inches(1.35),
        Inches(1.2),
        Inches(7.2),
        Inches(1.4),
        slide_data.get("title", "Academic Presentation"),
        font_name=TITLE_FONT,
        font_size=26,
        color=THEME["white"],
        bold=True,
    )
    add_text(
        slide,
        Inches(1.38),
        Inches(2.75),
        Inches(6.4),
        Inches(1.5),
        subtitle or "Course report / academic defense / research update",
        font_size=15,
        color="D6E2FF",
    )
    add_text(
        slide,
        Inches(1.38),
        Inches(5.75),
        Inches(4.2),
        Inches(0.4),
        ACTIVE_TEMPLATE.get("display_name", "Student PPT Pet"),
        font_size=11,
        color="B8CAE8",
    )
    add_text(
        slide,
        Inches(9.05),
        Inches(5.18),
        Inches(2.2),
        Inches(0.3),
        "Structured | Academic | Editable",
        font_size=10,
        color=THEME["white"],
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    return slide


def add_agenda_slide(prs: Presentation, slide_data: dict, page_num: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    variant = slide_data.get("variant", get_default_variant("agenda"))
    set_background(slide, THEME["light"])
    add_box(slide, Inches(0.45), Inches(0.45), Inches(12.25), Inches(0.55), THEME["navy"])
    add_text(slide, Inches(0.75), Inches(0.56), Inches(4), Inches(0.25), slide_data.get("title", "Agenda"), font_name=TITLE_FONT, font_size=24, color=THEME["white"], bold=True)
    items = normalize_lines(slide_data.get("content", []))
    if variant == "stacked-list":
        for idx, item in enumerate(items[:6], start=1):
            top = Inches(1.45) + (idx - 1) * Inches(0.88)
            add_box(slide, Inches(1.0), top, Inches(11.0), Inches(0.64), THEME["white"], THEME["border"])
            badge = add_box(slide, Inches(1.18), top + Inches(0.08), Inches(0.62), Inches(0.48), THEME["accent"])
            add_text(slide, badge.left, badge.top, badge.width, badge.height, f"{idx:02d}", font_name=TITLE_FONT, font_size=12, color=THEME["white"], bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
            add_text(slide, Inches(2.0), top + Inches(0.11), Inches(9.5), Inches(0.32), item, font_size=17, color=THEME["ink"], bold=True, valign=MSO_ANCHOR.MIDDLE)
        add_footer(slide, page_num)
        return slide

    positions = [
        (Inches(0.9), Inches(1.5)),
        (Inches(6.7), Inches(1.5)),
        (Inches(0.9), Inches(3.15)),
        (Inches(6.7), Inches(3.15)),
        (Inches(0.9), Inches(4.8)),
        (Inches(6.7), Inches(4.8)),
    ]
    for idx, item in enumerate(items[:6], start=1):
        left, top = positions[idx - 1]
        add_box(slide, left, top, Inches(5.0), Inches(1.2), THEME["white"], THEME["border"])
        circle = add_box(slide, left + Inches(0.18), top + Inches(0.16), Inches(0.75), Inches(0.75), THEME["blue"])
        add_text(
            slide,
            circle.left,
            circle.top + Inches(0.01),
            circle.width,
            circle.height,
            f"{idx:02d}",
            font_name=TITLE_FONT,
            font_size=15,
            color=THEME["white"],
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            left + Inches(1.1),
            top + Inches(0.18),
            Inches(3.55),
            Inches(0.7),
            item,
            font_name=DEFAULT_FONT,
            font_size=18,
            color=THEME["ink"],
            bold=True,
            valign=MSO_ANCHOR.MIDDLE,
        )
    add_footer(slide, page_num)
    return slide


def add_process_slide(prs: Presentation, slide_data: dict, page_num: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, THEME["white"])
    add_text(
        slide,
        Inches(0.7),
        Inches(0.55),
        Inches(8.5),
        Inches(0.5),
        slide_data.get("title", ""),
        font_name=TITLE_FONT,
        font_size=24,
        color=THEME["navy"],
        bold=True,
    )
    add_line(slide, Inches(0.72), Inches(1.1), Inches(2.6), Inches(0.05), THEME["accent"])

    steps = normalize_lines(slide_data.get("content", []))[:4] or [
        "Step 1",
        "Step 2",
        "Step 3",
        "Step 4",
    ]
    variant = slide_data.get("variant", get_default_variant("process"))
    if variant == "timeline-ribbon":
        add_line(slide, Inches(1.3), Inches(3.55), Inches(10.4), Inches(0.06), THEME["cyan"])
        for idx, step in enumerate(steps):
            left = Inches(1.2) + idx * Inches(2.7)
            badge = add_box(slide, left, Inches(2.85), Inches(0.72), Inches(0.72), THEME["accent"])
            add_text(slide, badge.left, badge.top + Inches(0.01), badge.width, badge.height, str(idx + 1), font_name=TITLE_FONT, font_size=15, color=THEME["white"], bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
            add_box(slide, left - Inches(0.15), Inches(4.0), Inches(1.8), Inches(1.18), THEME["sky"], THEME["border"])
            add_text(slide, left, Inches(4.22), Inches(1.45), Inches(0.7), step, font_size=14, color=THEME["ink"], bold=True, align=PP_ALIGN.CENTER)
        add_footer(slide, page_num)
        return slide

    step_width = Inches(2.75)
    step_height = Inches(2.35)
    base_top = Inches(2.0)
    for idx, step in enumerate(steps):
        left = Inches(0.75) + idx * Inches(3.05)
        add_box(slide, left, base_top, step_width, step_height, THEME["sky"], THEME["border"])
        badge = add_box(slide, left + Inches(0.22), base_top + Inches(0.2), Inches(0.52), Inches(0.52), THEME["blue"])
        add_text(
            slide,
            badge.left,
            badge.top,
            badge.width,
            badge.height,
            str(idx + 1),
            font_name=TITLE_FONT,
            font_size=14,
            color=THEME["white"],
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            left + Inches(0.22),
            base_top + Inches(0.9),
            Inches(2.2),
            Inches(1.0),
            step,
            font_size=17,
            color=THEME["ink"],
            bold=True,
        )
        if idx < len(steps) - 1:
            add_line(slide, left + Inches(2.8), base_top + Inches(1.12), Inches(0.32), Inches(0.04), THEME["cyan"])
    add_footer(slide, page_num)
    return slide


def add_results_chart(slide, chart_data_dict, left, top, width, height):
    categories = normalize_lines(chart_data_dict.get("categories", []))
    values = chart_data_dict.get("values", [])
    if not categories or not isinstance(values, list) or not values:
        return

    chart_data = ChartData()
    chart_data.categories = categories
    chart_data.add_series(chart_data_dict.get("series_name", "Series"), values[: len(categories)])

    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        left,
        top,
        width,
        height,
        chart_data,
    )
    chart = chart_shape.chart
    chart.has_legend = False
    chart.value_axis.has_major_gridlines = True
    chart.category_axis.tick_labels.font.size = Pt(11)
    chart.value_axis.tick_labels.font.size = Pt(11)
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.size = Pt(10)
    plot.series[0].format.fill.solid()
    plot.series[0].format.fill.fore_color.rgb = rgb(THEME["blue"])


def add_results_slide(prs: Presentation, slide_data: dict, page_num: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    variant = slide_data.get("variant", get_default_variant("results"))
    set_background(slide, "F5F8FE")
    add_text(
        slide,
        Inches(0.7),
        Inches(0.55),
        Inches(8.5),
        Inches(0.45),
        slide_data.get("title", ""),
        font_name=TITLE_FONT,
        font_size=24,
        color=THEME["navy"],
        bold=True,
    )
    add_line(slide, Inches(0.72), Inches(1.08), Inches(2.9), Inches(0.05), THEME["green"])

    if variant == "insight-grid":
        metrics = slide_data.get("metrics", [])
        for idx, metric in enumerate(metrics[:4]):
            left = Inches(0.9) + idx * Inches(2.95)
            add_box(slide, left, Inches(1.55), Inches(2.45), Inches(1.05), THEME["white"], THEME["border"])
            add_text(slide, left + Inches(0.12), Inches(1.74), Inches(2.15), Inches(0.28), str(metric.get("value", "--")), font_name=TITLE_FONT, font_size=20, color=THEME["navy"], bold=True, align=PP_ALIGN.CENTER)
            add_text(slide, left + Inches(0.12), Inches(2.18), Inches(2.15), Inches(0.2), str(metric.get("label", "")), font_size=10, color=THEME["gray"], align=PP_ALIGN.CENTER)
        insights = normalize_lines(slide_data.get("content", []))
        for idx, item in enumerate(insights[:4]):
            left = Inches(0.9) + (idx % 2) * Inches(5.8)
            top = Inches(3.05) + (idx // 2) * Inches(1.55)
            add_box(slide, left, top, Inches(5.15), Inches(1.15), THEME["white"], THEME["border"])
            add_text(slide, left + Inches(0.22), top + Inches(0.16), Inches(4.7), Inches(0.7), item, font_size=15, color=THEME["ink"], bold=True, valign=MSO_ANCHOR.MIDDLE)
        chart = slide_data.get("chart", {})
        if isinstance(chart, dict) and chart.get("categories") and chart.get("values"):
            add_results_chart(slide, chart, Inches(6.95), Inches(5.05), Inches(4.8), Inches(1.15))
        add_footer(slide, page_num)
        return slide

    add_box(slide, Inches(0.72), Inches(1.55), Inches(5.2), Inches(4.95), THEME["white"], THEME["border"])
    add_text(slide, Inches(1.0), Inches(1.85), Inches(4.5), Inches(0.3), "Key Findings", font_size=16, color=THEME["blue"], bold=True)
    add_bullets(slide, Inches(1.0), Inches(2.2), Inches(4.45), Inches(3.7), normalize_lines(slide_data.get("content", [])), font_size=17, color=THEME["ink"])

    metrics = slide_data.get("metrics", [])
    if isinstance(metrics, list) and metrics:
        for idx, metric in enumerate(metrics[:3]):
            left = Inches(6.2) + idx * Inches(2.05)
            add_box(slide, left, Inches(1.75), Inches(1.75), Inches(1.25), THEME["navy"])
            add_text(
                slide,
                left + Inches(0.12),
                Inches(1.95),
                Inches(1.5),
                Inches(0.35),
                str(metric.get("value", "--")),
                font_name=TITLE_FONT,
                font_size=20,
                color=THEME["white"],
                bold=True,
                align=PP_ALIGN.CENTER,
            )
            add_text(
                slide,
                left + Inches(0.1),
                Inches(2.55),
                Inches(1.55),
                Inches(0.35),
                str(metric.get("label", "")),
                font_size=11,
                color="DDE7FA",
                align=PP_ALIGN.CENTER,
            )
    else:
        cards = normalize_lines(slide_data.get("content", []))[:3]
        for idx, item in enumerate(cards):
            left = Inches(6.2)
            top = Inches(1.75) + idx * Inches(1.45)
            add_box(slide, left, top, Inches(5.4), Inches(1.15), THEME["white"], THEME["border"])
            add_line(slide, left, top, Inches(0.12), Inches(1.15), [THEME["blue"], THEME["green"], THEME["accent"]][idx % 3])
            add_text(
                slide,
                left + Inches(0.28),
                top + Inches(0.18),
                Inches(4.8),
                Inches(0.65),
                item,
                font_size=16,
                color=THEME["ink"],
                bold=True,
                valign=MSO_ANCHOR.MIDDLE,
            )

    chart = slide_data.get("chart", {})
    if isinstance(chart, dict) and chart.get("categories") and chart.get("values"):
        add_results_chart(slide, chart, Inches(6.3), Inches(4.55), Inches(5.2), Inches(1.5))

    add_footer(slide, page_num)
    return slide


def add_summary_slide(prs: Presentation, slide_data: dict, page_num: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    variant = slide_data.get("variant", get_default_variant("summary"))
    if variant == "closing-quote":
        set_background(slide, THEME["light"])
        add_box(slide, Inches(1.0), Inches(1.05), Inches(11.2), Inches(5.2), THEME["white"], THEME["border"])
        add_text(slide, Inches(1.55), Inches(1.55), Inches(10.0), Inches(0.5), slide_data.get("title", ""), font_name=TITLE_FONT, font_size=25, color=THEME["navy"], bold=True, align=PP_ALIGN.CENTER)
        quote = " / ".join(normalize_lines(slide_data.get("content", []))[:3]) or "Clear takeaway, sharp conclusion, confident delivery."
        add_text(slide, Inches(1.8), Inches(2.55), Inches(9.5), Inches(1.2), quote, font_size=20, color=THEME["ink"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_box(slide, Inches(4.55), Inches(4.75), Inches(4.15), Inches(0.72), THEME["accent"])
        add_text(slide, Inches(4.72), Inches(4.95), Inches(3.8), Inches(0.25), "Thanks • Q&A • Keep it sharp", font_size=13, color=THEME["white"], bold=True, align=PP_ALIGN.CENTER)
        add_footer(slide, page_num)
        return slide

    set_background(slide, THEME["navy"])
    add_text(
        slide,
        Inches(0.82),
        Inches(0.62),
        Inches(5.5),
        Inches(0.5),
        slide_data.get("title", ""),
        font_name=TITLE_FONT,
        font_size=24,
        color=THEME["white"],
        bold=True,
    )
    add_text(slide, Inches(0.82), Inches(1.32), Inches(4.6), Inches(0.35), "Key Takeaways", font_size=13, color="B8CAE8")

    add_box(slide, Inches(0.82), Inches(1.85), Inches(7.1), Inches(4.65), THEME["panel"])
    bullets = normalize_lines(slide_data.get("content", []))
    for idx, item in enumerate(bullets[:5]):
        top = Inches(2.15) + idx * Inches(0.82)
        badge = add_box(slide, Inches(1.12), top, Inches(0.32), Inches(0.32), THEME["accent"])
        add_text(
            slide,
            badge.left,
            badge.top - Inches(0.01),
            badge.width,
            badge.height,
            "+",
            font_size=11,
            color=THEME["white"],
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            Inches(1.58),
            top - Inches(0.02),
            Inches(5.7),
            Inches(0.42),
            item,
            font_size=17,
            color=THEME["white"],
            valign=MSO_ANCHOR.MIDDLE,
        )

    add_box(slide, Inches(8.35), Inches(1.85), Inches(3.95), Inches(4.65), THEME["panel_alt"])
    add_text(
        slide,
        Inches(8.7),
        Inches(2.2),
        Inches(3.2),
        Inches(0.4),
        "Defense Tips",
        font_name=TITLE_FONT,
        font_size=18,
        color=THEME["white"],
        bold=True,
    )
    hints = [
        "Start from the problem, then method, then evidence.",
        "Explain metrics in plain language instead of jargon.",
        "End with contribution and the next step.",
    ]
    add_bullets(slide, Inches(8.68), Inches(2.8), Inches(2.95), Inches(2.8), hints, font_size=15, color=THEME["white"])
    add_footer(slide, page_num)
    return slide


def add_content_slide(prs: Presentation, slide_data: dict, page_num: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    variant = slide_data.get("variant", get_default_variant("content"))
    set_background(slide, THEME["light"])
    if variant == "split-panel":
        add_text(slide, Inches(0.8), Inches(0.65), Inches(8.2), Inches(0.45), slide_data.get("title", ""), font_name=TITLE_FONT, font_size=24, color=THEME["navy"], bold=True)
        add_line(slide, Inches(0.82), Inches(1.15), Inches(2.7), Inches(0.05), THEME["accent"])
        add_box(slide, Inches(0.82), Inches(1.55), Inches(5.55), Inches(4.95), THEME["white"], THEME["border"])
        add_bullets(slide, Inches(1.1), Inches(1.95), Inches(4.95), Inches(4.0), normalize_lines(slide_data.get("content", [])), font_size=18, color=THEME["ink"])
        right_cards = normalize_lines(slide_data.get("content", []))[:3] or ["Key point", "Supporting point", "Takeaway"]
        for idx, item in enumerate(right_cards):
            top = Inches(1.7) + idx * Inches(1.52)
            add_box(slide, Inches(6.75), top, Inches(5.0), Inches(1.08), THEME["sky"], THEME["border"])
            add_text(slide, Inches(6.98), top + Inches(0.16), Inches(4.55), Inches(0.7), item, font_size=15, color=THEME["ink"], bold=True, valign=MSO_ANCHOR.MIDDLE)
        add_footer(slide, page_num)
        return slide

    add_box(slide, Inches(0.65), Inches(0.7), Inches(2.55), Inches(5.95), THEME["navy"])
    add_text(slide, Inches(0.92), Inches(1.2), Inches(2.0), Inches(1.5), slide_data.get("title", ""), font_name=TITLE_FONT, font_size=22, color=THEME["white"], bold=True)
    add_text(slide, Inches(0.95), Inches(5.7), Inches(1.8), Inches(0.35), "Academic Section", font_size=10, color="B8CAE8")
    add_box(slide, Inches(3.45), Inches(0.95), Inches(8.95), Inches(5.45), THEME["white"], THEME["border"])
    add_text(slide, Inches(3.82), Inches(1.28), Inches(3.5), Inches(0.28), "Core Content", font_size=15, color=THEME["blue"], bold=True)
    add_bullets(slide, Inches(3.82), Inches(1.7), Inches(5.65), Inches(4.2), normalize_lines(slide_data.get("content", [])), font_size=18, color=THEME["ink"])
    right_cards = normalize_lines(slide_data.get("content", []))[:2] or ["Point A", "Point B"]
    for idx, item in enumerate(right_cards):
        top = Inches(1.72) + idx * Inches(2.05)
        add_box(slide, Inches(9.75), top, Inches(2.25), Inches(1.55), THEME["sky"], THEME["border"])
        add_text(slide, Inches(9.96), top + Inches(0.18), Inches(1.8), Inches(1.1), item, font_size=14, color=THEME["ink"], bold=True, valign=MSO_ANCHOR.MIDDLE)

    add_footer(slide, page_num)
    return slide


def build_slide(prs: Presentation, slide_data: dict, page_num: int):
    layout = infer_layout(slide_data)
    if layout == "title":
        return add_cover_slide(prs, slide_data)
    if layout == "agenda":
        return add_agenda_slide(prs, slide_data, page_num)
    if layout == "process":
        return add_process_slide(prs, slide_data, page_num)
    if layout == "results":
        return add_results_slide(prs, slide_data, page_num)
    if layout in {"summary", "closing"}:
        return add_summary_slide(prs, slide_data, page_num)
    return add_content_slide(prs, slide_data, page_num)


def build_presentation(plan: dict) -> Presentation:
    normalized_plan = normalize_plan(plan, plan.get("template_key"))
    prs = Presentation()
    ensure_slide_size(prs)
    prs.core_properties.title = normalized_plan["title"]
    prs.core_properties.author = normalized_plan["author"]

    for idx, slide_data in enumerate(normalized_plan["slides"], start=1):
        build_slide(prs, slide_data, idx)
    return prs


def resolve_path(path_str: str | None, default: str | None = None) -> Path | None:
    value = path_str or default
    if not value:
        return None
    return Path(value).expanduser().resolve()


def merge_slide(existing_slide: dict, new_slide_data: dict, page_index: int) -> dict:
    if not isinstance(new_slide_data, dict):
        raise ValueError("new_data must be a JSON object for a single slide")

    merged = deepcopy(existing_slide)
    if "layout" in new_slide_data and "variant" not in new_slide_data:
        merged.pop("variant", None)
    for key in ("type", "layout", "variant", "title", "content", "metrics", "chart"):
        if key in new_slide_data:
            merged[key] = new_slide_data[key]
    if "notes" in new_slide_data:
        merged["notes"] = new_slide_data["notes"]
    return normalize_slide(merged, page_index)


def write_presentation(plan: dict, output_path: Path) -> str:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs = build_presentation(plan)
    prs.save(str(output_path))
    return str(output_path)


def create_presentation(
    plan_file: str,
    output_file: str,
    state_file: str | None,
    template_file: str | None,
    template_key: str | None,
) -> str:
    plan_path = resolve_path(plan_file)
    output_path = resolve_path(output_file, DEFAULT_OUTPUT)
    state_path = resolve_path(state_file, DEFAULT_STATE)
    template_path = resolve_path(template_file)
    raw_plan = load_json(plan_path)
    chosen_template_key = resolve_template_key(template_key, raw_plan)
    plan = normalize_plan(raw_plan, chosen_template_key)
    output_str = write_presentation(plan, output_path)
    save_json(
        state_path,
        make_state(
            plan,
            template_key=chosen_template_key,
            template_path=str(template_path) if template_path else "",
            output_path=str(output_path),
        ),
    )
    return output_str


def modify_presentation(
    *,
    state_file: str,
    page_index: int,
    new_data: str,
    output_file: str,
    template_file: str | None,
    template_key: str | None,
) -> str:
    state_path = resolve_path(state_file, DEFAULT_STATE)
    output_path = resolve_path(output_file, DEFAULT_OUTPUT)
    template_path = resolve_path(template_file)

    state = load_json(state_path)
    chosen_template_key = resolve_template_key(template_key, state)
    plan = normalize_plan(extract_plan_from_state(state), chosen_template_key)

    if page_index < 1 or page_index > len(plan["slides"]):
        raise ValueError(
            f"page_index out of range: {page_index}, total slides: {len(plan['slides'])}"
        )

    new_slide_data = parse_json_payload(new_data)
    slide_idx = page_index - 1
    plan["slides"][slide_idx] = merge_slide(plan["slides"][slide_idx], new_slide_data, page_index)
    output_str = write_presentation(plan, output_path)

    chosen_template = (
        str(template_path)
        if template_path
        else str(state.get("template_path", "")).strip()
    )
    save_json(
        state_path,
        make_state(
            plan,
            template_key=chosen_template_key,
            template_path=chosen_template,
            output_path=str(output_path),
        ),
    )
    return output_str


def restyle_presentation(
    *,
    state_file: str,
    output_file: str,
    template_file: str | None,
    template_key: str | None,
) -> str:
    state_path = resolve_path(state_file, DEFAULT_STATE)
    output_path = resolve_path(output_file, DEFAULT_OUTPUT)
    template_path = resolve_path(template_file)

    state = load_json(state_path)
    chosen_template_key = resolve_template_key(template_key, state)
    plan = restyle_plan_to_template_defaults(
        extract_plan_from_state(state), chosen_template_key
    )
    output_str = write_presentation(plan, output_path)

    chosen_template = (
        str(template_path)
        if template_path
        else str(state.get("template_path", "")).strip()
    )
    save_json(
        state_path,
        make_state(
            plan,
            template_key=chosen_template_key,
            template_path=chosen_template,
            output_path=str(output_path),
        ),
    )
    return output_str


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate or modify a styled PPTX file for student-ppt-pet"
    )
    parser.add_argument(
        "--action",
        choices=["create", "modify", "restyle"],
        default="create",
        help="Engine action. Use 'create' for full deck generation, 'modify' for single-page updates, or 'restyle' for template-only rerendering.",
    )
    parser.add_argument("--plan-file", "--plan_file", dest="plan_file", help="Path to JSON plan")
    parser.add_argument("--output-file", "--output", dest="output_file", default=DEFAULT_OUTPUT, help="Destination PPTX path")
    parser.add_argument("--state-file", "--state", dest="state_file", default=DEFAULT_STATE, help="State JSON path")
    parser.add_argument("--template-file", "--template", dest="template_file", help="Optional template path recorded in state metadata")
    parser.add_argument("--template-key", "--template_key", dest="template_key", help="Template key from assets/templates")
    parser.add_argument("--page-index", "--page_index", dest="page_index", type=int, help="1-based slide index for modify action")
    parser.add_argument("--new-data", "--new_data", dest="new_data", help="Single-slide JSON string for modify action")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    try:
        if args.action == "create":
            if not args.plan_file:
                raise ValueError("--plan-file is required for --action=create")
            output_path = create_presentation(
                args.plan_file, args.output_file, args.state_file, args.template_file, args.template_key
            )
        elif args.action == "modify":
            if not args.page_index:
                raise ValueError("--page-index is required for --action=modify")
            if not args.new_data:
                raise ValueError("--new-data is required for --action=modify")
            output_path = modify_presentation(
                state_file=args.state_file,
                page_index=args.page_index,
                new_data=args.new_data,
                output_file=args.output_file,
                template_file=args.template_file,
                template_key=args.template_key,
            )
        else:
            output_path = restyle_presentation(
                state_file=args.state_file,
                output_file=args.output_file,
                template_file=args.template_file,
                template_key=args.template_key,
            )
    except Exception as exc:
        print(f"[Error] {exc}")
        return 1

    print(f"[Success] Presentation successfully generated at: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
